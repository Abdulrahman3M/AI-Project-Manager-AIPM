#!/usr/bin/env python3
"""Build 00-Index.docx with relative hyperlinks to chapter DOCX files."""
from __future__ import annotations
import argparse, json, os
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_external_hyperlink(paragraph, text: str, target: str):
    part=paragraph.part
    rid=part.relate_to(target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hyperlink=OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"),rid)
    run=OxmlElement("w:r")
    rpr=OxmlElement("w:rPr")
    color=OxmlElement("w:color"); color.set(qn("w:val"),"0563C1")
    underline=OxmlElement("w:u"); underline.set(qn("w:val"),"single")
    rpr.append(color); rpr.append(underline)
    text_el=OxmlElement("w:t"); text_el.text=text
    run.append(rpr); run.append(text_el); hyperlink.append(run)
    paragraph._p.append(hyperlink)

def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("manifest")
    ap.add_argument("--output")
    ap.add_argument("--check-files",action="store_true")
    a=ap.parse_args()
    manifest_path=Path(a.manifest).resolve()
    data=json.loads(manifest_path.read_text(encoding="utf-8"))
    out=Path(a.output) if a.output else manifest_path.parent / data.get("index_file","00-Index.docx")
    if not out.is_absolute(): out=(manifest_path.parent/out).resolve()
    chapters=sorted(data.get("chapters",[]),key=lambda c:c["order"])
    if a.check_files:
        missing=[c["file"] for c in chapters if not (manifest_path.parent/c["file"]).is_file()]
        if missing:
            raise SystemExit("Missing chapter files: "+", ".join(missing))
    doc=Document()
    doc.add_heading(data.get("publication_title","Publication Index"),0)
    p=doc.add_paragraph("This index opens the separate Word chapter files. Keep the complete package in the same folder so relative links continue to work.")
    if data.get("description"): doc.add_paragraph(data["description"])
    doc.add_heading("Contents",1)
    for c in chapters:
        p=doc.add_paragraph()
        label=f'{c["order"]}. {c["title"]}'
        target=os.path.relpath((manifest_path.parent/c["file"]).resolve(),out.parent).replace(os.sep,"/")
        add_external_hyperlink(p,label,target)
        details=[]
        if c.get("estimated_pages"): details.append(f'~{c["estimated_pages"]} pages')
        if c.get("status"): details.append(str(c["status"]))
        if details: p.add_run(" — "+"; ".join(details))
        if c.get("description"): doc.add_paragraph(c["description"],style="Intense Quote")
    out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out)
    print(out)
    return 0

if __name__=="__main__":
    raise SystemExit(main())
